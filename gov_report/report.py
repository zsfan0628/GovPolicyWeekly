from __future__ import annotations

import json
import re
import shutil
import subprocess
import sys
from collections import Counter
from html import escape
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor




# Excel / Word 都不允许写入部分 XML 控制字符。
# 政府网站偶尔会把 .doc/.wps/.pdf 附件误当成正文抓回来，导致出现二进制乱码。
ILLEGAL_XML_RE = re.compile(r"[\x00-\x08\x0B\x0C\x0E-\x1F]")


def sanitize_export_value(value, max_len: int = 32000):
    """清理导出到 Excel/Word 的值，避免 openpyxl IllegalCharacterError。"""
    if value is None:
        return ""
    if isinstance(value, (list, tuple, set)):
        return "；".join(sanitize_export_value(v, max_len=4000) for v in value)
    if isinstance(value, dict):
        return json.dumps(sanitize_record(value), ensure_ascii=False)[:max_len]
    s = str(value)
    # 明显的二进制附件乱码：OLE 复合文档通常以 D0 CF 11 E0 开头，错误解码后常见 “��ࡱ”。
    if ("ࡱ" in s and s.count("�") >= 2) or s.count("�") >= 20:
        return "【疑似附件二进制乱码，已清理】"
    s = ILLEGAL_XML_RE.sub("", s)
    # Excel 单元格最大 32767 字符，留一点余量。
    if len(s) > max_len:
        s = s[:max_len] + "……【内容过长已截断】"
    return s


def sanitize_record(record: dict) -> dict:
    out = {}
    for k, v in (record or {}).items():
        if isinstance(v, dict):
            out[k] = sanitize_record(v)
        elif isinstance(v, list):
            out[k] = [sanitize_record(x) if isinstance(x, dict) else sanitize_export_value(x) for x in v]
        else:
            out[k] = sanitize_export_value(v)
    return out


REPORT_COLUMNS = [
    "publish_date",
    "source_name",
    "region",
    "department",
    "category",
    "related_directions",
    "title",
    "summary",
    "change_points",
    "url",
    "decision_source",
    "rule_reason",
    "in_policy_column",
    "official_document_type",
    "policy_subtype",
    "document_action",
    "llm_confidence",
    "llm_reason",
]


def export_all(records: list[dict], output_dir: str | Path, start, end, keywords: dict[str, list[str]]) -> dict[str, str]:
    records = [sanitize_record(r) for r in records]
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    basename = f"农药化工企业政策与人事任免周报_{start:%Y-%m-%d}_至_{end:%Y-%m-%d}"

    json_path = output_dir / "data.json"
    json_path.write_text(json.dumps(records, ensure_ascii=False, indent=2), encoding="utf-8")

    xlsx_path = output_dir / "data.xlsx"
    export_excel(records, xlsx_path)

    docx_path = output_dir / f"{basename}.docx"
    export_docx(records, docx_path, start, end, keywords)

    pdf_path = convert_docx_to_pdf(docx_path, output_dir)
    if not pdf_path:
        pdf_path = output_dir / f"{basename}.pdf"
        export_pdf_direct(records, pdf_path, start, end, keywords)

    return {
        "json": str(json_path),
        "xlsx": str(xlsx_path),
        "docx": str(docx_path),
        "pdf": str(pdf_path) if pdf_path else "",
    }


def export_excel(records: list[dict], path: str | Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    path = Path(path)
    columns = REPORT_COLUMNS + ["attachments"]
    wb = Workbook()
    ws = wb.active
    ws.title = "data"
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    header_font = Font(bold=True)

    ws.append(columns)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for r in records:
        row = []
        for key in REPORT_COLUMNS:
            row.append(r.get(key, ""))
        row.append(
            "; ".join(
                f"{a.get('name','附件')} {a.get('url','')} {a.get('local_path','')}"
                for a in r.get("attachments", [])
            )
        )
        ws.append(row)

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    ws.freeze_panes = "A2"
    for idx, column in enumerate(columns, 1):
        width = 18
        if column in {"title", "summary", "change_points", "url", "attachments", "llm_reason"}:
            width = 36
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = width

    wb.save(path)


def export_docx(records: list[dict], path: str | Path, start, end, keywords: dict[str, list[str]]) -> None:
    path = Path(path)
    doc = Document()
    setup_document(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("农药化工企业政策与人事任免信息周报")
    run.bold = True
    run.font.size = Pt(20)
    set_run_font(run)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run(f"统计周期：{start:%Y年%m月%d日} 至 {end:%Y年%m月%d日}")
    r.font.size = Pt(11)
    set_run_font(r)

    gen = doc.add_paragraph()
    gen.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = gen.add_run("说明：本报告由规则硬通过 + GPT-5.5复判生成；政策栏目中的通知、函、决定、公告、公示等正式文种按政策文件收录；模糊候选交由GPT复判；项目环评受理/拟审查/审批决定等单个项目公示默认不按政策文件收录。")
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(100, 100, 100)
    set_run_font(r)

    doc.add_heading("一、本周概览", level=1)
    add_overview(doc, records)

    doc.add_heading("二、信息清单", level=1)
    add_list_table(doc, records)

    doc.add_heading("三、详细条目", level=1)
    if not records:
        p = doc.add_paragraph("本周期内未筛选到符合关键词规则的条目。")
        set_paragraph_font(p)
    else:
        for i, record in enumerate(records, 1):
            add_detail_record(doc, i, record, keywords)

    doc.save(path)


def setup_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)
    styles = doc.styles
    styles["Normal"].font.name = "宋体"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    styles["Normal"].font.size = Pt(10.5)
    for style_name in ["Heading 1", "Heading 2", "Heading 3"]:
        style = styles[style_name]
        style.font.name = "黑体"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")


def set_run_font(run, font_name: str = "宋体") -> None:
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)


def set_paragraph_font(p, font_name: str = "宋体") -> None:
    for run in p.runs:
        set_run_font(run, font_name)


def add_overview(doc: Document, records: list[dict]) -> None:
    by_category = Counter()
    by_direction = Counter()
    by_dept = Counter()
    for r in records:
        for c in str(r.get("category", "")).split("、"):
            if c:
                by_category[c] += 1
        for d in str(r.get("related_directions", "")).split("、"):
            if d:
                by_direction[d] += 1
        if r.get("department"):
            by_dept[r["department"]] += 1

    rows = [
        ("命中总数", len(records)),
        ("人事任免", by_category.get("人事任免", 0)),
        ("政策文件", by_category.get("政策文件", 0)),
        ("涉及农药/农资", by_direction.get("农药/农资", 0)),
        ("涉及化工/危化品", by_direction.get("化工/危化品", 0)),
        ("涉及生态环境/排污", by_direction.get("生态环境/排污", 0)),
        ("涉及安全生产", by_direction.get("安全生产", 0)),
        ("涉及产业扶持/项目申报", by_direction.get("产业扶持/项目申报", 0)),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "项目"
    hdr[1].text = "数量"
    shade_cell(hdr[0], "D9EAF7")
    shade_cell(hdr[1], "D9EAF7")
    for k, v in rows:
        cells = table.add_row().cells
        cells[0].text = str(k)
        cells[1].text = str(v)
    format_table(table)


def add_list_table(doc: Document, records: list[dict]) -> None:
    cols = ["日期", "发布单位", "类型", "相关方向", "标题", "判定方式", "变动/附件"]
    table = doc.add_table(rows=1, cols=len(cols))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, col in enumerate(cols):
        table.rows[0].cells[i].text = col
        shade_cell(table.rows[0].cells[i], "D9EAF7")

    for r in records:
        cells = table.add_row().cells
        cells[0].text = r.get("publish_date", "")
        cells[1].text = r.get("source_name", "")
        cells[2].text = r.get("category", "")
        cells[3].text = r.get("related_directions", "")
        title_paragraph = cells[4].paragraphs[0]
        add_hyperlink(title_paragraph, r.get("title", ""), r.get("url", ""))
        flags = []
        if r.get("change_points"):
            flags.append("有变动提示")
        if r.get("attachments"):
            flags.append(f"附件{len(r.get('attachments'))}个")
        cells[5].text = r.get("decision_source", "")
        cells[6].text = "；".join(flags)
    format_table(table, font_size=8.0)


def add_detail_record(doc: Document, idx: int, record: dict, keywords: dict[str, list[str]]) -> None:
    h = doc.add_heading(f"{idx}. {record.get('title','')}", level=2)
    set_paragraph_font(h, "黑体")

    meta = doc.add_paragraph()
    meta.add_run("发布时间：").bold = True
    meta.add_run(record.get("publish_date", ""))
    meta.add_run("    发布单位：").bold = True
    meta.add_run(record.get("source_name", ""))
    meta.add_run("    类型：").bold = True
    meta.add_run(record.get("category", ""))
    set_paragraph_font(meta)

    p = doc.add_paragraph()
    p.add_run("相关方向：").bold = True
    p.add_run(record.get("related_directions", ""))
    set_paragraph_font(p)

    p = doc.add_paragraph()
    p.add_run("判定方式：").bold = True
    p.add_run(record.get("decision_source", ""))
    if record.get("official_document_type"):
        p.add_run("    文种：").bold = True
        p.add_run(str(record.get("official_document_type", "")))
    if record.get("policy_subtype"):
        p.add_run("    政策子类：").bold = True
        p.add_run(str(record.get("policy_subtype", "")))
    if record.get("llm_confidence") not in (None, ""):
        p.add_run("    GPT置信度：").bold = True
        p.add_run(str(record.get("llm_confidence", "")))
    set_paragraph_font(p)

    if record.get("llm_reason"):
        p = doc.add_paragraph()
        p.add_run("GPT判断理由：").bold = True
        p.add_run(record.get("llm_reason", ""))
        set_paragraph_font(p)

    p = doc.add_paragraph()
    p.add_run("原文链接：").bold = True
    add_hyperlink(p, record.get("url", ""), record.get("url", ""))

    doc.add_paragraph("精简摘要：")
    p = doc.add_paragraph()
    add_highlighted_text(p, record.get("summary", ""), keywords)

    if record.get("change_points"):
        doc.add_paragraph("变动/调整提示：")
        for point in str(record.get("change_points", "")).splitlines():
            p = doc.add_paragraph(style=None)
            p.paragraph_format.left_indent = Inches(0.25)
            p.add_run("• ")
            add_highlighted_text(p, point, keywords)

    attachments = record.get("attachments") or []
    if attachments:
        doc.add_paragraph("附件：")
        for att in attachments:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.add_run("• ")
            name = att.get("name") or "附件"
            url = att.get("url") or ""
            local_path = att.get("local_path") or ""
            add_hyperlink(p, name, url)
            if local_path:
                p.add_run(f"    本地：{local_path}")
            set_paragraph_font(p)


def add_highlighted_text(p, text: str, keywords: dict[str, list[str]]) -> None:
    if not text:
        return
    # 优先级：橙色取消/废止，蓝色支持/补贴，紫色监管，红色一般调整。
    color_rules = []
    for k in keywords.get("orange_cancel", []):
        color_rules.append((k, RGBColor(230, 126, 34), True))
    for k in keywords.get("blue_support", []):
        color_rules.append((k, RGBColor(41, 128, 185), True))
    for k in keywords.get("purple_regulation", []):
        color_rules.append((k, RGBColor(125, 60, 152), True))
    for k in keywords.get("change", []):
        color_rules.append((k, RGBColor(192, 0, 0), True))

    pos = 0
    while pos < len(text):
        match = None
        for k, color, bold in color_rules:
            if k and text.startswith(k, pos):
                match = (k, color, bold)
                break
        if match:
            k, color, bold = match
            run = p.add_run(k)
            run.font.color.rgb = color
            run.bold = bold
            set_run_font(run)
            pos += len(k)
        else:
            run = p.add_run(text[pos])
            set_run_font(run)
            pos += 1


def add_hyperlink(paragraph, text: str, url: str):
    if not text:
        return
    if not url:
        run = paragraph.add_run(text)
        set_run_font(run)
        return
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    rPr.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(u)
    new_run.append(rPr)
    t = OxmlElement("w:t")
    t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def format_table(table, font_size: float = 9.0) -> None:
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    set_run_font(run)
                    run.font.size = Pt(font_size)


def convert_docx_to_pdf(docx_path: str | Path, out_dir: str | Path) -> Path | None:
    """使用 LibreOffice 转 PDF。没有安装时返回 None。"""
    docx_path = Path(docx_path)
    out_dir = Path(out_dir)
    app_dir = Path(sys.executable).resolve().parent if getattr(sys, "frozen", False) else Path.cwd()
    candidates = [
        shutil.which("soffice"),
        shutil.which("libreoffice"),
        app_dir / "LibreOffice" / "program" / "soffice.exe",
        app_dir / "_internal" / "LibreOffice" / "program" / "soffice.exe",
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
        "/Applications/LibreOffice.app/Contents/MacOS/soffice",
    ]
    soffice = next((c for c in candidates if c and Path(c).exists()), None)
    if not soffice:
        return None
    cmd = [
        soffice,
        "--headless",
        "--convert-to",
        "pdf",
        "--outdir",
        str(out_dir),
        str(docx_path),
    ]
    try:
        subprocess.run(cmd, check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True, timeout=120)
    except Exception:
        return None
    pdf_path = out_dir / (docx_path.stem + ".pdf")
    return pdf_path if pdf_path.exists() else None


def export_pdf_direct(records: list[dict], path: str | Path, start, end, keywords: dict[str, list[str]]) -> Path:
    """使用 reportlab 直接生成 PDF，避免依赖 LibreOffice。"""
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))

    doc = SimpleDocTemplate(
        str(path),
        pagesize=A4,
        leftMargin=16 * mm,
        rightMargin=16 * mm,
        topMargin=16 * mm,
        bottomMargin=16 * mm,
        title="农药化工企业政策与人事任免信息周报",
    )
    styles = getSampleStyleSheet()
    base = ParagraphStyle(
        "CNBase",
        parent=styles["Normal"],
        fontName="STSong-Light",
        fontSize=9.5,
        leading=14,
        wordWrap="CJK",
        spaceAfter=5,
    )
    title_style = ParagraphStyle(
        "CNTitle",
        parent=base,
        fontSize=18,
        leading=24,
        alignment=TA_CENTER,
        spaceAfter=8,
    )
    subtitle_style = ParagraphStyle(
        "CNSubtitle",
        parent=base,
        fontSize=10,
        leading=14,
        alignment=TA_CENTER,
        textColor=colors.HexColor("#555555"),
        spaceAfter=10,
    )
    heading = ParagraphStyle("CNHeading", parent=base, fontSize=13, leading=18, spaceBefore=8, spaceAfter=6)
    item_title = ParagraphStyle("CNItemTitle", parent=base, fontSize=11, leading=16, spaceBefore=8, spaceAfter=4)
    small = ParagraphStyle("CNSmall", parent=base, fontSize=8, leading=11)

    story = [
        Paragraph("农药化工企业政策与人事任免信息周报", title_style),
        Paragraph(f"统计周期：{start:%Y年%m月%d日} 至 {end:%Y年%m月%d日}", subtitle_style),
        Paragraph(
            "说明：本报告由规则硬通过 + GPT复判生成；政策栏目中的通知、函、决定、公告、公示等正式文种按政策文件收录；模糊候选交由GPT复判。",
            small,
        ),
        Spacer(1, 5),
        Paragraph("一、本周概览", heading),
    ]

    overview_rows = _pdf_overview_rows(records)
    overview_table = Table([[Paragraph(str(a), base), Paragraph(str(b), base)] for a, b in overview_rows], colWidths=[65 * mm, 95 * mm])
    overview_table.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (-1, -1), "STSong-Light"),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#b8c7d6")),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#d9eaf7")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    story.append(overview_table)
    story.append(Spacer(1, 8))
    story.append(Paragraph("二、信息清单", heading))

    list_rows = [[
        Paragraph("日期", small),
        Paragraph("发布单位", small),
        Paragraph("类型", small),
        Paragraph("标题", small),
    ]]
    for r in records:
        list_rows.append(
            [
                Paragraph(_pdf_text(r.get("publish_date", "")), small),
                Paragraph(_pdf_text(r.get("source_name", "")), small),
                Paragraph(_pdf_text(r.get("category", "")), small),
                Paragraph(_pdf_link(r.get("title", ""), r.get("url", "")), small),
            ]
        )
    if len(list_rows) == 1:
        list_rows.append([Paragraph("本周期内未筛选到符合关键词规则的条目。", base), "", "", ""])
    list_table = Table(list_rows, colWidths=[23 * mm, 37 * mm, 25 * mm, 78 * mm], repeatRows=1)
    list_table.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (-1, -1), "STSong-Light"),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#c7c7c7")),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#d9eaf7")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    story.append(list_table)
    story.append(PageBreak())
    story.append(Paragraph("三、详细条目", heading))

    if not records:
        story.append(Paragraph("本周期内未筛选到符合关键词规则的条目。", base))
    else:
        for idx, record in enumerate(records, 1):
            story.append(Paragraph(f"{idx}. {_pdf_text(record.get('title', ''))}", item_title))
            story.append(Paragraph(f"发布时间：{_pdf_text(record.get('publish_date', ''))}    发布单位：{_pdf_text(record.get('source_name', ''))}    类型：{_pdf_text(record.get('category', ''))}", base))
            story.append(Paragraph(f"相关方向：{_pdf_text(record.get('related_directions', ''))}", base))
            story.append(Paragraph(f"判定方式：{_pdf_text(record.get('decision_source', ''))}", base))
            if record.get("llm_reason"):
                story.append(Paragraph(f"GPT判断理由：{_pdf_text(record.get('llm_reason', ''))}", base))
            if record.get("summary"):
                story.append(Paragraph(f"精简摘要：{_pdf_text(record.get('summary', ''))}", base))
            if record.get("change_points"):
                story.append(Paragraph(f"变动/调整提示：{_pdf_text(record.get('change_points', ''))}", base))
            if record.get("url"):
                story.append(Paragraph(f"原文链接：{_pdf_text(record.get('url', ''))}", small))
            attachments = record.get("attachments") or []
            if attachments:
                att_text = "；".join([_pdf_text(a.get("name") or "附件") for a in attachments])
                story.append(Paragraph(f"附件：{att_text}", small))
            story.append(Spacer(1, 5))

    doc.build(story)
    return path


def _pdf_overview_rows(records: list[dict]) -> list[tuple[str, int]]:
    by_category = Counter()
    by_direction = Counter()
    for r in records:
        for c in str(r.get("category", "")).split("、"):
            if c:
                by_category[c] += 1
        for d in str(r.get("related_directions", "")).split("、"):
            if d:
                by_direction[d] += 1
    return [
        ("命中总数", len(records)),
        ("人事任免", by_category.get("人事任免", 0)),
        ("政策文件", by_category.get("政策文件", 0)),
        ("涉及农药/农资", by_direction.get("农药/农资", 0)),
        ("涉及化工/危化品", by_direction.get("化工/危化品", 0)),
        ("涉及生态环境/排污", by_direction.get("生态环境/排污", 0)),
        ("涉及安全生产", by_direction.get("安全生产", 0)),
        ("涉及产业扶持/项目申报", by_direction.get("产业扶持/项目申报", 0)),
    ]


def _pdf_text(value) -> str:
    text = sanitize_export_value(value, max_len=4000)
    return (
        str(text)
        .replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace("\n", "<br/>")
    )


def _pdf_attr(value) -> str:
    return escape(str(value or ""), quote=True)


def _pdf_link(text, url) -> str:
    label = _pdf_text(text)
    href = str(url or "").strip()
    if not href:
        return label
    return f'<a href="{_pdf_attr(href)}">{label}</a>'
